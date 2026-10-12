from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def apply_dark_theme(slide):
    # Set dark background
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(20, 24, 35) # Deep navy/gray background
    
    # Add a cyan accent line at the top
    left = Inches(0)
    top = Inches(0)
    width = Inches(10)
    height = Inches(0.15)
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0, 200, 255) # Cyan accent
    shape.line.color.rgb = RGBColor(0, 200, 255)

def format_text(shape, font_size, is_title=False):
    if not shape.has_text_frame:
        return
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(font_size)
            run.font.name = 'Segoe UI'
            if is_title:
                run.font.color.rgb = RGBColor(0, 200, 255) # Cyan for titles
                run.font.bold = True
            else:
                run.font.color.rgb = RGBColor(230, 230, 230) # Light gray for text

def add_slide_with_content(prs, title_text, bullet_points):
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    apply_dark_theme(slide)
    
    title = slide.shapes.title
    title.text = title_text
    format_text(title, 36, is_title=True)
    
    content = slide.placeholders[1]
    tf = content.text_frame
    
    for i, point in enumerate(bullet_points):
        if i == 0:
            tf.text = point
        else:
            p = tf.add_paragraph()
            p.text = point
            p.level = 0
            
    format_text(content, 22, is_title=False)
    return slide

def create_presentation():
    prs = Presentation()
    # Ensure slide size is 16:9 (Widescreen)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # -----------------------------------------
    # Slide 1: Title Slide
    # -----------------------------------------
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    apply_dark_theme(slide)
    
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "SIH 2026: AI Email Threat Intelligence"
    subtitle.text = "Project Update & Review Presentation\nPrepared for Mentor Review"
    
    format_text(title, 54, is_title=True)
    format_text(subtitle, 28, is_title=False)
    
    # -----------------------------------------
    # Slide 2: Project Overview & Context
    # -----------------------------------------
    add_slide_with_content(prs, "Project Overview & Context", [
        "Domain: Cyber Security / Threat Intelligence",
        "Goal: Build an automated, Tier-1 SOC analyst pipeline for email threats.",
        "Focus: Deep forensic analysis on raw .eml files rather than simple text filtering.",
        "Target Audience: Security Operations Centers (SOC) and enterprise security teams."
    ])
    
    # -----------------------------------------
    # Slide 3: Problem Statement
    # -----------------------------------------
    add_slide_with_content(prs, "The Problem", [
        "Traditional spam filters are easily bypassed by sophisticated phishing and spoofing attacks.",
        "Attackers leverage complex infrastructure (Tor, VPNs, spoofed headers) to hide their origins.",
        "Manual forensic analysis of .eml files by security analysts is time-consuming and error-prone.",
        "Need a unified system to parse, verify, and score threats automatically."
    ])

    # -----------------------------------------
    # Slide 4: Objectives & Milestones
    # -----------------------------------------
    add_slide_with_content(prs, "Key Objectives", [
        "1. Automate metadata parsing and network hop reconstruction.",
        "2. Verify security authentications (SPF, DKIM, DMARC).",
        "3. Implement Machine Learning for content and NLP-based anomaly detection.",
        "4. Map IP origins to flag suspicious infrastructure (VPNs, Tor nodes).",
        "5. Develop a React-based Dashboard for actionable insights."
    ])

    # -----------------------------------------
    # Slide 5: System Architecture
    # -----------------------------------------
    add_slide_with_content(prs, "System Architecture & Tech Stack", [
        "Backend: Python (FastAPI) for high-performance API processing.",
        "Frontend: React.js for the SOC Analyst Dashboard.",
        "Database: SQLite/MySQL managed via SQLAlchemy ORM.",
        "Machine Learning: Scikit-learn for NLP and classification.",
        "Network Analysis: dnspython (DNS lookups), geoip2 (IP geolocation)."
    ])

    # -----------------------------------------
    # Slide 6: Implementation - The Pipeline
    # -----------------------------------------
    add_slide_with_content(prs, "The Core Pipeline", [
        "Phase 1: Ingestion & Parsing (Extracting headers, body, attachments)",
        "Phase 2: Hop Reconstruction (Tracing the 'Received' header path)",
        "Phase 3: Security Validation (Checking cryptographic signatures)",
        "Phase 4: ML Classification (Scoring text for phishing probability)",
        "Phase 5: Geolocation (Extracting true origin IP & mapping infrastructure)"
    ])

    # -----------------------------------------
    # Slide 7: Current Progress & Results
    # -----------------------------------------
    add_slide_with_content(prs, "Current Progress & Results", [
        "Backend API (/analyze endpoint) is fully functional and processing .eml files.",
        "Database schema (EmailCase) is successfully persisting forensic data.",
        "Composite Fraud Score logic is implemented, combining ML and hard security failures.",
        "Basic test files successfully parse and return accurate JSON intelligence reports."
    ])

    # -----------------------------------------
    # Slide 8: Challenges & Roadblocks
    # -----------------------------------------
    add_slide_with_content(prs, "Challenges & Solutions", [
        "Challenge: Accurately reconstructing spoofed hop paths.",
        "Solution: Implemented strict validation of the earliest reliable hop.",
        "Challenge: Handling varying .eml formats and encodings.",
        "Solution: Utilized robust Python email parsing libraries with fallback decoding.",
        "Challenge: Real-time ML processing latency.",
        "Solution: Optimized the pipeline by using efficient Scikit-learn models."
    ])

    # -----------------------------------------
    # Slide 9: Next Steps & Future Work
    # -----------------------------------------
    add_slide_with_content(prs, "Next Steps", [
        "1. Complete the React.js frontend SOC Dashboard integration.",
        "2. Improve ML model accuracy with a larger dataset of malicious emails.",
        "3. Implement attachment scanning and sandboxing.",
        "4. Prepare API webhooks for SIEM (e.g., Splunk) integration."
    ])

    # -----------------------------------------
    # Slide 10: Feedback & Discussion
    # -----------------------------------------
    add_slide_with_content(prs, "Feedback & Discussion", [
        "Open for mentor feedback on:",
        "- Architecture choices and scalability.",
        "- Fraud score weighting logic.",
        "- Additional features for the SIH 2026 final submission.",
        "Thank You!"
    ])

    prs.save('SIH2026_Mentor_Presentation_Styled.pptx')
    print("Presentation saved successfully as SIH2026_Mentor_Presentation_Styled.pptx")

if __name__ == '__main__':
    create_presentation()
